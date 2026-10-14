"""Batch questionnaires for ``lens render --batch`` (CSV via stdlib, XLSX via optional openpyxl).

Input columns pass through unchanged; ``answer``, ``citations``, ``confidence`` and
``stale_sources`` are appended. Processing is sequential and resumable: when the output file
already exists, rows whose ``answer`` is filled are kept as they are.
"""

from __future__ import annotations

import csv
import os
import re
import sys
from collections.abc import Callable, Sequence
from pathlib import Path
from typing import IO, Any

from gtm_telescope.render import RenderResult

OUTPUT_COLUMNS = ("answer", "citations", "confidence", "stale_sources")
_QUESTION_HEADER_RE = re.compile(r"question|requirement|query", re.IGNORECASE)
XLSX_SUFFIXES = frozenset({".xlsx", ".xlsm"})

Row = list[Any]


class BatchError(RuntimeError):
    """User-facing batch failure (the CLI prints it and exits 1)."""


def _require_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as e:
        raise BatchError(
            "XLSX support needs openpyxl. Install the batch extra: pip install '.[batch]'"
        ) from e
    return openpyxl


def _is_xlsx(path: Path) -> bool:
    return path.suffix.lower() in XLSX_SUFFIXES


def _cell(v: Any) -> str:
    return "" if v is None else str(v)


def read_table(path: Path, sheet: str | None = None) -> tuple[list[str], list[Row]]:
    """Return ``(headers, rows)``. Cells keep their native values for XLSX."""
    if not path.is_file():
        raise BatchError(f"batch file not found: {path}")
    if _is_xlsx(path):
        openpyxl = _require_openpyxl()
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            if sheet is not None:
                if sheet not in wb.sheetnames:
                    raise BatchError(f"sheet {sheet!r} not found in {path} (have: {wb.sheetnames})")
                ws = wb[sheet]
            else:
                ws = wb.active
            all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            all_rows = [list(r) for r in csv.reader(f)]
    # Drop fully empty trailing rows (spreadsheets often carry them).
    while all_rows and not any(_cell(c).strip() for c in all_rows[-1]):
        all_rows.pop()
    if not all_rows:
        raise BatchError(f"no rows found in {path}")
    headers = [_cell(h).strip() for h in all_rows[0]]
    width = len(headers)
    rows = [(r + [""] * width)[:width] for r in all_rows[1:]]
    return headers, rows


def pick_question_column(headers: Sequence[str], requested: str | None) -> int:
    """``--question-col`` (case-insensitive), else first header matching question/requirement/query,
    else the first column."""
    if requested is not None:
        for i, h in enumerate(headers):
            if h.strip().lower() == requested.strip().lower():
                return i
        raise BatchError(f"question column {requested!r} not found (headers: {list(headers)})")
    for i, h in enumerate(headers):
        if _QUESTION_HEADER_RE.search(h):
            return i
    return 0


def output_fields(result: RenderResult) -> list[str]:
    return [
        result.answer,
        "; ".join(c.location for c in result.citations),
        result.confidence,
        "; ".join(dict.fromkeys(c.path for c in result.citations if c.stale)),
    ]


def default_out_path(batch_path: Path) -> Path:
    return batch_path.with_name(f"{batch_path.stem}.answers{batch_path.suffix}")


def _write_atomic(
    path: Path, headers: Sequence[str], rows: Sequence[Row], sheet: str | None
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        if _is_xlsx(path):
            openpyxl = _require_openpyxl()
            wb = openpyxl.Workbook()
            ws = wb.active
            if sheet:
                ws.title = sheet[:31]
            ws.append(list(headers))
            for r in rows:
                ws.append(list(r))
            wb.save(str(tmp))
        else:
            with tmp.open("w", encoding="utf-8", newline="") as f:
                w = csv.writer(f)
                w.writerow(headers)
                w.writerows(rows)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def _load_resume_rows(
    out_path: Path, headers: Sequence[str], n_rows: int, sheet: str | None
) -> dict[int, Row]:
    """Rows of an existing output file whose ``answer`` is filled, keyed by row index."""
    if not out_path.is_file():
        return {}
    out_headers, out_rows = read_table(out_path, sheet)
    expected = [*headers, *OUTPUT_COLUMNS]
    if out_headers != expected or len(out_rows) != n_rows:
        raise BatchError(
            f"refusing to resume: {out_path} does not match the input file "
            "(different headers or row count). Remove it or choose another --out."
        )
    answer_idx = len(headers)
    return {i: r for i, r in enumerate(out_rows) if _cell(r[answer_idx]).strip()}


def run_batch(
    batch_path: Path,
    *,
    answer: Callable[[str], RenderResult],
    question_col: str | None = None,
    out_path: Path | None = None,
    sheet: str | None = None,
    progress: IO[str] | None = None,
) -> Path:
    """Answer every question row. Returns the output path.

    ``answer`` maps a question string to a :class:`RenderResult`. Output is rewritten atomically
    after every row so an interrupted run can be resumed by re-running the same command.
    """
    err = progress if progress is not None else sys.stderr
    headers, rows = read_table(batch_path, sheet)
    q_idx = pick_question_column(headers, question_col)
    out = out_path or default_out_path(batch_path)
    out_headers = [*headers, *OUTPUT_COLUMNS]
    done = _load_resume_rows(out, headers, len(rows), sheet)
    total = len(rows)
    out_rows: list[Row] = []
    for i, row in enumerate(rows):
        out_rows.append(done.get(i, [*row, *([""] * len(OUTPUT_COLUMNS))]))

    for i, row in enumerate(rows):
        question = _cell(row[q_idx]).strip()
        if i in done:
            err.write(f"[{i + 1}/{total}] (resumed) {question[:70]}\n")
            continue
        if not question:
            err.write(f"[{i + 1}/{total}] (blank question, passed through)\n")
            continue
        err.write(f"[{i + 1}/{total}] {question[:70]}\n")
        err.flush()
        result = answer(question)
        out_rows[i] = [*row, *output_fields(result)]
        _write_atomic(out, out_headers, out_rows, sheet)
    _write_atomic(out, out_headers, out_rows, sheet)  # also covers all-resumed / all-blank runs
    return out


__all__ = [
    "OUTPUT_COLUMNS",
    "BatchError",
    "default_out_path",
    "output_fields",
    "pick_question_column",
    "read_table",
    "run_batch",
]
